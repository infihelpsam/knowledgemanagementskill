import csv
import hashlib
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from bs4 import BeautifulSoup


def normalize(text: str):
    text = unicodedata.normalize('NFKC', text)
    text = ''.join(c for c in text if c in '\n\t' or not unicodedata.category(c).startswith('C'))
    return re.sub(r'[ \t]+', ' ', text).strip()


def extract(data: bytes, filename: str, maximum: int, max_chars: int):
    if len(data) > maximum:
        raise ValueError('File exceeds size limit')
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 2000:
            raise ValueError('Encrypted or excessive-page PDF')
        parts = []
        total = 0
        for page in reader.pages:
            text = page.extract_text() or ''
            total += len(text)
            if total > max_chars:
                raise ValueError('Extracted text exceeds limit')
            parts.append(text)
        text = '\n'.join(parts)
    elif suffix == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > maximum * 10:
                raise ValueError('DOCX decompression limit exceeded')
        doc = Document(io.BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
    elif suffix in {'.txt', '.md', '.markdown', '.html', '.htm', '.csv'}:
        text = data.decode('utf-8-sig', errors='strict')
        if suffix in {'.html', '.htm'}:
            soup = BeautifulSoup(text, 'html.parser')
            for tag in soup(['script', 'style', 'noscript', 'iframe', 'object']):
                tag.decompose()
            text = soup.get_text('\n')
        elif suffix == '.csv':
            text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(text)))
    else:
        raise ValueError('Unsupported extension')
    if len(text) > max_chars:
        raise ValueError('Extracted text exceeds limit')
    text = normalize(text)
    if not text:
        raise ValueError('Document has no extractable text; OCR is not included')
    return text


def fingerprint(text):
    return hashlib.sha256(text.encode()).hexdigest()


def chunks(text, size=1600, overlap=200):
    if not 0 <= overlap < size:
        raise ValueError('Invalid chunk overlap')
    for index, start in enumerate(range(0, len(text), size - overlap)):
        yield {'ordinal': index, 'start': start, 'end': min(start + size, len(text)), 'text': text[start:start + size]}
        if start + size >= len(text):
            break
